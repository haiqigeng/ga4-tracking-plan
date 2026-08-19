from __future__ import annotations

import argparse
import importlib.util
import json
import os
from pathlib import Path
from typing import Any
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from template_fidelity import drawing_object_snapshot

NATIVE_REQUIRED_PREFIXES = {
    "chart": "Excel-authored charts and drawing extensions are not proven safe through the ordinary writer.",
    "drawing": "Drawing relationships and anchors require native Excel preservation.",
    "image": "Image anchors and drawing relationships require native Excel preservation.",
    "pivot": "Pivot tables and caches are not safely editable through the ordinary writer.",
    "slicer": "Slicers and slicer caches are not supported by the ordinary writer.",
    "external_link": "External-link packages require native Excel with link updates disabled.",
    "custom_xml": "Custom XML is not modelled by the ordinary writer.",
    "macro": "Macro-enabled workbooks require native Excel with macros disabled during adaptation.",
    "connection": "Workbook data connections are not proven safe through the ordinary writer.",
    "query_table": "Query tables and their refresh metadata require native Excel preservation.",
    "metadata": "Excel metadata and rich-value packages are not modelled by the ordinary writer.",
    "threaded_comment": "Threaded comments and person metadata require native Excel preservation.",
    "web_extension": "Office web extensions are not modelled by the ordinary writer.",
    "custom_ui": "Custom ribbon/UI packages require native Office preservation.",
    "timeline": "Pivot timelines and caches are not supported by the ordinary writer.",
    "control": "Worksheet control properties require native Excel preservation.",
    "printer_setting": "Binary printer settings are not modelled by the ordinary writer.",
    "data_model": "Workbook data-model parts require native Excel preservation.",
}

BLOCKING_PREFIXES = {
    "encrypted": "Encrypted workbooks cannot be inspected or adapted safely.",
    "activex": "ActiveX controls are an unverified mutation surface.",
    "ole_object": "Embedded OLE objects are an unverified mutation surface.",
    "digital_signature": "Digitally signed workbook packages cannot be mutated without invalidating the signature.",
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Inventory a supplied workbook and select a fidelity-safe writer.")
    parser.add_argument("template", type=Path)
    parser.add_argument("--output", "-o", type=Path, required=True)
    return parser.parse_args()


def native_excel_available() -> bool:
    if os.name != "nt":
        return False
    try:
        return importlib.util.find_spec("win32com.client") is not None
    except (ImportError, ModuleNotFoundError, ValueError):
        return False


def _package_parts(path: Path) -> list[str]:
    try:
        with ZipFile(path) as archive:
            return sorted(name for name in archive.namelist() if not name.endswith("/"))
    except (BadZipFile, OSError):
        return []


def _feature(feature: str, location: str, reason: str) -> dict[str, str]:
    return {"feature": feature, "location": location, "reason": reason}


def _formula_inventory(workbook: Any) -> tuple[list[dict[str, Any]], int]:
    samples: list[dict[str, Any]] = []
    count = 0
    for sheet in workbook.worksheets:
        for cell in sheet._cells.values():
            if not (isinstance(cell.value, str) and cell.value.startswith("=")):
                continue
            count += 1
            if len(samples) < 100:
                samples.append({"sheet": sheet.title, "coordinate": cell.coordinate, "formula": cell.value})
    return samples, count


def _defined_name_inventory(workbook: Any) -> list[dict[str, Any]]:
    values: list[dict[str, Any]] = []
    for name in workbook.defined_names.values():
        values.append(
            {
                "name": str(name.name),
                "attr_text": str(name.attr_text or ""),
                "local_sheet_id": name.localSheetId,
                "hidden": bool(name.hidden),
            }
        )
    return sorted(values, key=lambda item: (item["name"], item["local_sheet_id"] if item["local_sheet_id"] is not None else -1))


def _sheet_inventory(sheet: Any) -> dict[str, Any]:
    tables = []
    for table in sheet.tables.values():
        tables.append(
            {
                "name": str(table.name),
                "display_name": str(table.displayName),
                "ref": str(table.ref),
                "totals_row_shown": bool(table.totalsRowShown),
                "style": str(getattr(getattr(table, "tableStyleInfo", None), "name", "") or ""),
            }
        )
    validations = [
        {
            "sqref": str(item.sqref),
            "type": str(item.type or ""),
            "formula1": str(item.formula1 or ""),
            "formula2": str(item.formula2 or ""),
        }
        for item in sheet.data_validations.dataValidation
    ]
    conditional_formats = [
        {
            "sqref": str(conditional.sqref),
            "rule_types": [str(rule.type or "") for rule in rules],
            "rule_count": len(rules),
        }
        for conditional, rules in sheet.conditional_formatting._cf_rules.items()
    ]
    pivots = [str(getattr(pivot, "name", "") or getattr(pivot, "cacheId", "")) for pivot in getattr(sheet, "_pivots", [])]
    return {
        "title": sheet.title,
        "state": sheet.sheet_state,
        "dimensions": sheet.calculate_dimension(),
        "tables": sorted(tables, key=lambda item: item["name"]),
        "data_validations": sorted(validations, key=lambda item: item["sqref"]),
        "conditional_formatting": sorted(conditional_formats, key=lambda item: item["sqref"]),
        "chart_count": len(sheet._charts),
        "image_count": len(sheet._images),
        "charts": [
            drawing_object_snapshot(item, include_model=False)
            for item in sheet._charts
        ],
        "images": [
            drawing_object_snapshot(item, include_model=False)
            for item in sheet._images
        ],
        "pivot_count": len(pivots),
        "pivots": pivots,
        "merged_ranges": sorted(str(item) for item in sheet.merged_cells.ranges),
        "freeze_panes": str(sheet.freeze_panes or ""),
        "auto_filter": str(sheet.auto_filter.ref or ""),
        "print_area": str(sheet.print_area or ""),
        "print_title_rows": str(sheet.print_title_rows or ""),
        "print_title_cols": str(sheet.print_title_cols or ""),
        "sheet_view_count": len(sheet.views.sheetView),
    }


def inspect_template_richness(path: Path) -> dict[str, Any]:
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("Supplied templates must be XLSX or XLSM workbooks.")
    parts = _package_parts(path)
    if not parts:
        raise ValueError("The supplied template is not a readable OOXML workbook package.")
    workbook = load_workbook(
        path,
        data_only=False,
        read_only=False,
        keep_links=True,
        keep_vba=path.suffix.lower() == ".xlsm",
    )
    formulas, formula_count = _formula_inventory(workbook)
    sheets = [_sheet_inventory(sheet) for sheet in workbook.worksheets]
    custom_xml_parts = [name for name in parts if name.startswith("customXml/")]
    external_link_parts = [name for name in parts if name.startswith("xl/externalLinks/")]
    pivot_parts = [name for name in parts if name.startswith("xl/pivot")]
    slicer_parts = [name for name in parts if name.startswith("xl/slicer") or "slicer" in name.casefold()]
    drawing_parts = [name for name in parts if name.startswith("xl/drawings/")]
    chart_parts = [name for name in parts if name.startswith("xl/charts/")]
    image_parts = [name for name in parts if name.startswith("xl/media/")]
    active_x_parts = [name for name in parts if name.startswith("xl/activeX/")]
    ole_parts = [name for name in parts if name.startswith("xl/embeddings/")]
    signature_parts = [name for name in parts if "signature" in name.casefold()]
    macro_parts = [name for name in parts if name.casefold().endswith("vbaproject.bin")]
    connection_parts = [name for name in parts if name.casefold() == "xl/connections.xml"]
    query_table_parts = [name for name in parts if name.casefold().startswith("xl/querytables/")]
    metadata_parts = [
        name
        for name in parts
        if name.casefold() == "xl/metadata.xml"
        or name.casefold().startswith("xl/richdata/")
    ]
    threaded_comment_parts = [
        name
        for name in parts
        if name.casefold().startswith(("xl/threadedcomments/", "xl/persons/"))
    ]
    web_extension_parts = [name for name in parts if name.casefold().startswith("xl/webextensions/")]
    custom_ui_parts = [name for name in parts if name.casefold().startswith("customui/")]
    timeline_parts = [name for name in parts if name.casefold().startswith("xl/timeline")]
    control_parts = [name for name in parts if name.casefold().startswith("xl/ctrlprops/")]
    printer_setting_parts = [name for name in parts if name.casefold().startswith("xl/printersettings/")]
    data_model_parts = [name for name in parts if name.casefold().startswith("xl/model/")]

    native_required: list[dict[str, str]] = []
    blocking: list[dict[str, str]] = []

    def add_many(target: list[dict[str, str]], prefix: str, locations: list[str]) -> None:
        reason = (NATIVE_REQUIRED_PREFIXES | BLOCKING_PREFIXES)[prefix]
        for location in locations:
            target.append(_feature(prefix, location, reason))

    add_many(native_required, "chart", chart_parts)
    add_many(native_required, "drawing", [name for name in drawing_parts if name not in chart_parts])
    add_many(native_required, "image", image_parts)
    add_many(native_required, "pivot", pivot_parts)
    add_many(native_required, "slicer", slicer_parts)
    add_many(native_required, "external_link", external_link_parts)
    add_many(native_required, "custom_xml", custom_xml_parts)
    add_many(native_required, "macro", macro_parts or ([path.name] if path.suffix.lower() == ".xlsm" else []))
    add_many(native_required, "connection", connection_parts)
    add_many(native_required, "query_table", query_table_parts)
    add_many(native_required, "metadata", metadata_parts)
    add_many(native_required, "threaded_comment", threaded_comment_parts)
    add_many(native_required, "web_extension", web_extension_parts)
    add_many(native_required, "custom_ui", custom_ui_parts)
    add_many(native_required, "timeline", timeline_parts)
    add_many(native_required, "control", control_parts)
    add_many(native_required, "printer_setting", printer_setting_parts)
    add_many(native_required, "data_model", data_model_parts)
    add_many(blocking, "activex", active_x_parts)
    add_many(blocking, "ole_object", ole_parts)
    add_many(blocking, "digital_signature", signature_parts)

    native_available = native_excel_available()
    if blocking:
        decision = "blocked"
        recommended_writer = "blocked"
    elif native_required:
        decision = "native_required"
        recommended_writer = "native_excel" if native_available else "blocked"
        if not native_available:
            blocking.append(
                _feature(
                    "native_excel_unavailable",
                    str(path),
                    "This template requires native Excel preservation, but the native adapter is unavailable. Install Excel plus pywin32 or use an owner-approved simplified template.",
                )
            )
            decision = "blocked"
    else:
        decision = "supported"
        recommended_writer = "openpyxl"

    return {
        "inventory_version": "1.0.0",
        "path": str(path.resolve()),
        "extension": path.suffix.lower(),
        "package_part_count": len(parts),
        "package_parts": parts,
        "sheets": sheets,
        "formula_count": formula_count,
        "formula_samples": formulas,
        "defined_names": _defined_name_inventory(workbook),
        "workbook_properties": {
            "date1904": bool(getattr(workbook.epoch, "year", 1899) == 1904),
            "calculation_mode": str(getattr(workbook.calculation, "calcMode", "") or ""),
            "full_calculation_on_load": bool(getattr(workbook.calculation, "fullCalcOnLoad", False)),
            "force_full_calculation": bool(getattr(workbook.calculation, "forceFullCalc", False)),
        },
        "native_excel_available": native_available,
        "native_required_features": native_required,
        "blocking_features": blocking,
        "decision": decision,
        "recommended_writer": recommended_writer,
    }


def ensure_writer_supported(preflight: dict[str, Any], requested: str = "auto") -> str:
    if requested not in {"auto", "openpyxl", "native_excel"}:
        raise ValueError(f"Unknown supplied-template writer: {requested}")
    if preflight.get("blocking_features"):
        details = "; ".join(
            f"{item.get('feature')} at {item.get('location')}: {item.get('reason')}" for item in preflight["blocking_features"]
        )
        raise ValueError("Supplied-template richness preflight blocked adaptation. " + details)
    recommended = str(preflight.get("recommended_writer", "blocked"))
    selected = recommended if requested == "auto" else requested
    if selected == "openpyxl" and preflight.get("native_required_features"):
        names = sorted({str(item.get("feature")) for item in preflight["native_required_features"]})
        raise ValueError("The ordinary writer is not proven safe for: " + ", ".join(names) + ". Use the native Excel adapter.")
    if selected == "native_excel" and not preflight.get("native_excel_available"):
        raise ValueError("The native Excel adapter requires Windows, Microsoft Excel, and pywin32.")
    if selected == "blocked":
        raise ValueError("No fidelity-safe writer is available for this supplied template.")
    return selected


def main() -> int:
    args = parse_args()
    try:
        result = inspect_template_richness(args.template)
    except (OSError, ValueError) as error:
        print(str(error))
        return 2
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(args.output)
    return 0 if result["decision"] != "blocked" else 1


if __name__ == "__main__":
    raise SystemExit(main())
